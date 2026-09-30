"""
Concatenate Paper A's section files into one readable draft.

The manuscript is written as fragments so that sections can be revised and
frozen independently -- §2 has been frozen since rev. 4 -- but a reader needs one
document. This joins them in order, lifts each section's drafting notes out of
the body and collects them at the end, and appends captions and references.

    python3 analysis/assemble_draft.py          # -> manuscript/PaperA_draft.md
    python3 analysis/assemble_draft.py --docx   # also a .docx

Nothing here edits content. If a number is wrong it is wrong in the section file.
"""
import re
import sys
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
M = ROOT / "manuscript"
OUT_MD = M / "PaperA_draft.md"

TITLE = ("Temporary Loss of Eligibility and Bias in "
         "Cross-Sectional HIV Incidence Estimation")
AUTHOR = "A. C. Demidont, Nyx Dynamics LLC. ORCID 0000-0002-9216-8569."

ORDER = [
    ("section0_abstract.md",        None),
    ("section1_introduction.md",    None),
    ("section2_theory.md",          None),
    ("section3_parameterisation.md", None),
    ("section4_results.md",         None),
    ("section5_discussion.md",      None),
    ("section6_supplement.md",      "Supplement"),
    ("captions.md",                 "Figures and tables"),
    ("references.md",               "References"),
]

NOTE_HEAD = re.compile(r"^#+\s*Drafting notes.*$", re.M)

# Blocks that belong to a section file but not to a flowing draft: the abstract's
# own title line (the document already has one) and its structured alternate,
# which is a variant for other venues rather than part of the manuscript.
LIFT = {
    "section0_abstract.md": [r"^##\s*Title\s*$", r"^##\s*Structured variant\s*$"],
}


def lift_blocks(text: str, patterns):
    """Remove a '## Heading' and everything under it up to the next '## '."""
    for pat in patterns:
        m = re.search(pat, text, re.M)
        if not m:
            continue
        nxt = re.search(r"^##\s+\S", text[m.end():], re.M)
        end = m.end() + (nxt.start() if nxt else len(text[m.end():]))
        text = text[:m.start()] + text[end:]
    return text


def demote(text: str) -> str:
    """Push '## ' to '### ' so reference sub-lists do not read as sections."""
    return re.sub(r"^##\s+", "### ", text, flags=re.M)


def split_notes(text: str):
    """(body, notes). Drafting notes are editorial and must not read as content."""
    m = NOTE_HEAD.search(text)
    if not m:
        return text.strip(), ""
    return text[:m.start()].strip(), text[m.end():].strip()


def strip_preamble(text: str) -> str:
    """Drop the per-file '# Paper A — Section n' banner and any leading rule."""
    lines = text.split("\n")
    out, dropped = [], False
    for ln in lines:
        if not dropped and (ln.startswith("# Paper A") or ln.startswith("> ")):
            continue
        if not dropped and ln.strip() in ("", "---"):
            continue
        dropped = True
        out.append(ln)
    return "\n".join(out).strip()


def main():
    parts, notes = [], []
    parts.append(f"# {TITLE}\n\n*{AUTHOR}*\n")
    parts.append(f"> Assembled draft. Regenerate with `python3 "
                 f"analysis/assemble_draft.py`; edit the section files, not this.\n")

    for fname, heading in ORDER:
        p = M / fname
        if not p.exists():
            print(f"  missing {fname}, skipped")
            continue
        body, note = split_notes(p.read_text())
        if fname in LIFT:
            body = lift_blocks(body, LIFT[fname])
        body = strip_preamble(body)
        if heading:
            body = demote(body)
        if heading:
            parts.append(f"---\n\n## {heading}\n")
        parts.append(body)
        if note:
            notes.append(f"### From `{fname}`\n\n{note}")

    if notes:
        parts.append("---\n\n# Drafting notes\n\n*Editorial. Not part of the "
                     "manuscript; collected here so the body reads clean.*\n")
        parts.extend(notes)

    text = "\n\n".join(parts) + "\n"
    OUT_MD.write_text(text)

    body_only = text.split("\n# Drafting notes")[0]
    print(f"  wrote {OUT_MD.relative_to(ROOT)}")
    print(f"  {len(body_only.split()):,} words of manuscript, "
          f"{len(text.split()) - len(body_only.split()):,} words of drafting notes")

    if "--docx" in sys.argv:
        try:
            from docx import Document
            from docx.shared import Pt, Inches
        except ImportError:
            print("  python-docx not installed; skipped .docx")
            return 0
        doc = Document()
        st = doc.styles["Normal"]
        st.font.name, st.font.size = "Times New Roman", Pt(11)
        for sec in doc.sections:
            sec.left_margin = sec.right_margin = Inches(1.0)
        for ln in text.split("\n"):
            s = ln.rstrip()
            if not s:
                continue
            lvl = len(s) - len(s.lstrip("#"))
            if lvl:
                doc.add_heading(s.lstrip("# ").strip(), min(lvl, 4))
            else:
                doc.add_paragraph(re.sub(r"[*_`]", "", s))
        out = M / "PaperA_draft.docx"
        doc.save(out)
        print(f"  wrote {out.relative_to(ROOT)}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
